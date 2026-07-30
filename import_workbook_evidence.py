"""Import completed workbook evidence layers into the website SQLite database."""

from __future__ import annotations

import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK_PATH = Path(
    "/Users/jjoseph/Desktop/geo_scrap_2026_april28th/"
    "geo_metadata_with_ai_FINAL_COMPLETED_CLEAN.xlsx"
)


def _split_gses(value):
    if value is None:
        return []
    text = str(value).replace(";", ",")
    return [part.strip() for part in text.split(",") if part.strip()]


def _cell(value):
    if value is None:
        return ""
    return str(value).strip()


def _headers(sheet):
    return [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]


def import_evidence(db_path: Path, workbook_path: Path = WORKBOOK_PATH):
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    if not db_path.exists():
        raise FileNotFoundError(db_path)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    conn = sqlite3.connect(db_path)

    conn.executescript(
        """
        DROP TABLE IF EXISTS evidence_rows;
        DROP TABLE IF EXISTS supplement_repair_rows;

        CREATE TABLE evidence_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            gse_id TEXT NOT NULL,
            paper_key TEXT,
            pmcid TEXT,
            doi TEXT,
            title TEXT,
            evidence_source_type TEXT,
            model TEXT,
            question TEXT,
            answer TEXT,
            confidence TEXT,
            quote TEXT,
            source TEXT,
            reason TEXT
        );

        CREATE INDEX idx_evidence_rows_gse ON evidence_rows(gse_id);

        CREATE TABLE supplement_repair_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            gse_id TEXT NOT NULL,
            paper_key TEXT,
            pmcid TEXT,
            doi TEXT,
            title TEXT,
            question TEXT,
            paper_answer TEXT,
            supplement_answer TEXT,
            merge_action TEXT,
            confidence TEXT,
            reason TEXT,
            quote TEXT,
            source TEXT,
            chunk_id TEXT,
            repair_file TEXT
        );

        CREATE INDEX idx_supplement_repair_rows_gse ON supplement_repair_rows(gse_id);
        """
    )

    evidence_count = 0
    evidence_sheet = workbook["Evidence"]
    evidence_headers = _headers(evidence_sheet)
    e = {name: index for index, name in enumerate(evidence_headers)}
    for row in evidence_sheet.iter_rows(min_row=2, values_only=True):
        for gse_id in _split_gses(row[e["Linked GEO Series ID(s)"]]):
            conn.execute(
                """
                INSERT INTO evidence_rows (
                    gse_id, paper_key, pmcid, doi, title, evidence_source_type,
                    model, question, answer, confidence, quote, source, reason
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    gse_id,
                    _cell(row[e["PaperKey"]]),
                    _cell(row[e["Linked PMCID(s)"]]),
                    _cell(row[e["Linked DOI(s)"]]),
                    _cell(row[e["Linked Title(s)"]]),
                    _cell(row[e["Evidence Source Type"]]),
                    _cell(row[e["Model"]]),
                    _cell(row[e["Question"]]),
                    _cell(row[e["Answer"]]),
                    _cell(row[e["Confidence"]]),
                    _cell(row[e["Quote"]]),
                    _cell(row[e["Source"]]),
                    _cell(row[e["Reason"]]),
                ],
            )
            evidence_count += 1

    repair_count = 0
    repair_sheet = workbook["Supplement_Repair_Findings"]
    repair_headers = _headers(repair_sheet)
    r = {name: index for index, name in enumerate(repair_headers)}
    for row in repair_sheet.iter_rows(min_row=2, values_only=True):
        for gse_id in _split_gses(row[r["Linked GEO Series ID(s)"]]):
            conn.execute(
                """
                INSERT INTO supplement_repair_rows (
                    gse_id, paper_key, pmcid, doi, title, question, paper_answer,
                    supplement_answer, merge_action, confidence, reason, quote,
                    source, chunk_id, repair_file
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    gse_id,
                    _cell(row[r["PaperKey"]]),
                    _cell(row[r["Linked PMCID(s)"]]),
                    _cell(row[r["Linked DOI(s)"]]),
                    _cell(row[r["Linked Title(s)"]]),
                    _cell(row[r["Question"]]),
                    _cell(row[r["PaperAnswer"]]),
                    _cell(row[r["SupplementAnswer"]]),
                    _cell(row[r["MergeAction"]]),
                    _cell(row[r["Confidence"]]),
                    _cell(row[r["Reason"]]),
                    _cell(row[r["Quote"]]),
                    _cell(row[r["Source"]]),
                    _cell(row[r["ChunkID"]]),
                    _cell(row[r["_RepairFile"]]),
                ],
            )
            repair_count += 1

    conn.commit()
    conn.close()
    return evidence_count, repair_count


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: import_workbook_evidence.py /path/to/geo_metadata.db")
    evidence_rows, repair_rows = import_evidence(Path(sys.argv[1]))
    print(f"evidence_rows={evidence_rows}")
    print(f"supplement_repair_rows={repair_rows}")
